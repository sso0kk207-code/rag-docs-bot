"""Загрузка документов: извлечение текста (txt, md, pdf, docx) и нарезка на фрагменты."""
import io
import re
from dataclasses import dataclass

MAX_CHARS = 800      # целевой размер фрагмента
OVERLAP_SENTENCES = 1  # перекрытие: последняя фраза предыдущего фрагмента повторяется в следующем, чтобы мысль не рвалась на границе


class UnsupportedFormat(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


@dataclass
class Chunk:
    text: str
    page: int | None


def extract_pages(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Возвращает [(номер страницы | None, текст)]."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("txt", "md", "markdown"):
        return [(None, data.decode("utf-8", errors="replace"))]
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [(i, page.extract_text() or "") for i, page in enumerate(reader.pages, start=1)]
    if ext == "docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:  # таблицы тоже содержат факты
            lines += [" | ".join(c.text for c in row.cells) for row in table.rows]
        return [(None, "\n\n".join(lines))]
    raise UnsupportedFormat(f"Unsupported file type: .{ext or '?'} (supported: txt, md, pdf, docx)")


SENTENCE = re.compile(r"(?<=[.!?])\s+")


def _sections(text: str) -> list[tuple[str, str]]:
    """Делит markdown по заголовкам: [(заголовок, тело)]. Заголовок попадёт в каждый фрагмент раздела как контекст."""
    parts, heading, body = [], "", []
    for line in text.splitlines():
        if re.match(r"^#{1,4}\s+\S", line):
            if "".join(body).strip():
                parts.append((heading, "\n".join(body)))
            heading, body = line.lstrip("# ").strip(), []
        else:
            body.append(line)
    if "".join(body).strip():
        parts.append((heading, "\n".join(body)))
    return parts


def chunk_text(text: str, page: int | None = None, max_chars: int = MAX_CHARS) -> list[Chunk]:
    chunks: list[Chunk] = []
    for heading, body in _sections(text):
        sentences: list[str] = []
        for para in re.split(r"\n\s*\n", body):
            para = " ".join(para.split())
            if para:
                sentences += SENTENCE.split(para)
        current: list[str] = []
        size = 0
        for s in sentences:
            if current and size + len(s) > max_chars:
                chunks.append(Chunk(_compose(heading, current), page))
                current = current[-OVERLAP_SENTENCES:]
                size = sum(len(x) for x in current)
            current.append(s)
            size += len(s) + 1
        if current:
            chunks.append(Chunk(_compose(heading, current), page))
    return chunks


def _compose(heading: str, sentences: list[str]) -> str:
    body = " ".join(sentences).strip()
    return f"{heading}. {body}" if heading else body


def ingest(filename: str, data: bytes) -> tuple[list[Chunk], int]:
    pages = extract_pages(filename, data)
    chunks = [c for page, text in pages for c in chunk_text(text, page)]
    if not chunks:
        raise EmptyDocument("No text could be extracted from the file")
    return chunks, len(pages)

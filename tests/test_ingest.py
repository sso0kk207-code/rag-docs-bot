import io

import docx

from app.rag.ingest import MAX_CHARS, chunk_text
from tests.conftest import upload
from tests.pdfgen import make_pdf


def test_markdown_is_split_by_headings_and_keeps_heading_as_context():
    text = "# Отпуск\n\n## Перенос\nДни можно перенести. Но не более 14.\n\n## Больничный\nСообщите руководителю."
    chunks = chunk_text(text)
    assert [c.text.split(".")[0] for c in chunks] == ["Перенос", "Больничный"]
    assert "Дни можно перенести" in chunks[0].text and "руководителю" in chunks[1].text


def test_long_section_is_split_with_overlap_and_size_limit():
    sentences = [f"Предложение номер {i} содержит достаточно слов для проверки." for i in range(60)]
    chunks = chunk_text("## Раздел\n" + " ".join(sentences))
    assert len(chunks) > 3
    assert all(len(c.text) <= MAX_CHARS + 200 for c in chunks)
    # перекрытие: последняя фраза предыдущего фрагмента открывает следующий
    last_of_first = chunks[0].text.split(". ")[-1].split(".")[0]
    assert last_of_first in chunks[1].text


def test_chunks_are_never_empty():
    assert chunk_text("\n\n\n") == []
    assert all(c.text.strip() for c in chunk_text("# A\n\n# B\ntext\n\n\n# C\n"))


def test_upload_markdown(client):
    r = upload(client, "policy.md", "# Заголовок\nПервое предложение. Второе предложение.".encode())
    assert r.status_code == 201 and r.json()["chunks"] == 1 and r.json()["pages"] == 1
    assert [d["filename"] for d in client.get("/documents").json()] == ["policy.md"]


def test_upload_pdf_keeps_page_numbers(client):
    pdf = make_pdf(["Vacation policy: twenty eight days.", "Expense policy: per diem is 3500."])
    r = upload(client, "rules.pdf", pdf, "application/pdf")
    assert r.status_code == 201 and r.json()["pages"] == 2
    hits = client.post("/search", json={"query": "per diem expense"}).json()
    assert hits[0]["page"] == 2 and "3500" in hits[0]["text"]


def test_upload_docx_including_tables(client):
    d = docx.Document()
    d.add_paragraph("Правила хранения данных клиентов.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Срок хранения", "пять лет"
    buf = io.BytesIO()
    d.save(buf)
    assert upload(client, "rules.docx", buf.getvalue()).status_code == 201
    assert "пять лет" in client.post("/search", json={"query": "срок хранения"}).json()[0]["text"]


def test_rejections(client):
    assert upload(client, "image.png", b"\x89PNG").status_code == 415
    assert upload(client, "empty.md", b"   \n\n  ").status_code == 422
    assert upload(client, "broken.pdf", b"not really a pdf").status_code == 422
    assert upload(client, "huge.txt", b"x" * (5 * 1024 * 1024 + 1)).status_code == 413


def test_same_file_is_not_uploaded_twice(client):
    first = upload(client, "a.md", "Один и тот же текст документа.".encode())
    second = upload(client, "renamed.md", "Один и тот же текст документа.".encode())
    assert first.status_code == 201 and second.status_code == 409 and str(first.json()["id"]) in second.json()["detail"]


def test_delete_document_removes_it_from_search(client):
    doc = upload(client, "secret.md", "# Секрет\nКод от сейфа 4815162342.".encode()).json()
    assert client.post("/search", json={"query": "код от сейфа"}).json()
    assert client.delete(f"/documents/{doc['id']}").status_code == 204
    assert client.post("/search", json={"query": "код от сейфа"}).json() == []
    assert client.delete(f"/documents/{doc['id']}").status_code == 404
